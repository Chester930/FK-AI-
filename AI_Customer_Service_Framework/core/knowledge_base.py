from config import KNOWLEDGE_BASE_PATH
import os
import re
import docx
import openpyxl
import PyPDF2
import fitz  # PyMuPDF
import jieba

class KnowledgeBase:
    def __init__(self, folderpath=KNOWLEDGE_BASE_PATH):
        self.folderpath = folderpath
        self.data = self.load_data()

    def load_data(self):
        """Loads data from multiple file types in the knowledge base folder."""
        all_text = ""
        for filename in os.listdir(self.folderpath):
            filepath = os.path.join(self.folderpath, filename)
            if filename.endswith(".txt"):
                all_text += self.read_txt(filepath) + "\n"
            elif filename.endswith(".docx"):
                all_text += self.read_docx(filepath) + "\n"
            elif filename.endswith(".xlsx"):
                all_text += self.read_xlsx(filepath) + "\n"
            elif filename.endswith(".pdf"):
                all_text += self.read_pdf(filepath) + "\n"
            elif os.path.isdir(filepath):
                all_text += self.read_folder(filepath) + "\n"
        return all_text

    def read_txt(self, filepath):
        """Reads text from a .txt file."""
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                return f.read()
        except Exception as e:
            print(f"Error reading TXT file {filepath}: {e}")
            return ""

    def read_docx(self, filepath):
        """Reads text from a .docx file."""
        try:
            doc = docx.Document(filepath)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            return "\n".join(full_text)
        except Exception as e:
            print(f"Error reading DOCX file {filepath}: {e}")
            return ""

    def read_xlsx(self, filepath):
        """Reads text from a .xlsx file."""
        try:
            workbook = openpyxl.load_workbook(filepath)
            all_text = []
            for sheet in workbook:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value:
                            all_text.append(str(cell.value))
            return "\n".join(all_text)
        except Exception as e:
            print(f"Error reading XLSX file {filepath}: {e}")
            return ""

    def read_pdf(self, filepath):
        """Reads text from a PDF file."""
        try:
            all_text = []
            # 使用 PyMuPDF
            with fitz.open(filepath) as doc:
                for page in doc:
                    all_text.append(page.get_text())
            return "\n".join(all_text)
        except Exception as e:
            print(f"Error reading PDF file {filepath}: {e}")
            return ""

    def read_folder(self, folderpath):
        """Recursively reads text from all files in a folder."""
        all_text = ""
        for filename in os.listdir(folderpath):
            filepath = os.path.join(folderpath, filename)
            if filename.endswith(".txt"):
                all_text += self.read_txt(filepath) + "\n"
            elif filename.endswith(".docx"):
                all_text += self.read_docx(filepath) + "\n"
            elif filename.endswith(".xlsx"):
                all_text += self.read_xlsx(filepath) + "\n"
            elif filename.endswith(".pdf"):
                all_text += self.read_pdf(filepath) + "\n"
            elif os.path.isdir(filepath):
                all_text += self.read_folder(filepath) + "\n"
        return all_text

    def search(self, query):
        """Searches the knowledge base for a query.

        Args:
            query (str): The search query.

        Returns:
            str: A relevant portion of the knowledge base, or a message
                 indicating that no relevant information was found.
        """
        if not self.data:
            return "Knowledge base is empty."

        # 使用 jieba 進行分詞
        query_words = jieba.lcut(query)

        # 使用布林搜尋 (AND 邏輯)
        matched_lines = []
        for line in self.data.split("\n"):
            if all(word.lower() in line.lower() for word in query_words):
                matched_lines.append(line)

        if matched_lines:
            # 只返回前 3 個匹配項
            return "\n".join(matched_lines[:3])
        else:
            return "No relevant information found in the knowledge base."